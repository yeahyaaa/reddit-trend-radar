import openpyxl
import pytest

from radar.models import Post
from radar.xlsx import write_xlsx

RANKED = [
    Post(
        title="Ubuntu 26.04 released",
        subreddit="ubuntu",
        score=900,
        num_comments=120,
        permalink="/r/ubuntu/comments/1",
    ),
    Post(
        title="Snap vs flatpak",
        subreddit="ubuntu",
        permalink="/r/ubuntu/comments/2",
        source="rss",
    ),
]


@pytest.fixture
def sheet(tmp_path):
    target = tmp_path / "out.xlsx"
    write_xlsx(RANKED, target)
    return openpyxl.load_workbook(target).active


def cell(sheet, row, heading):
    """Look a value up by column heading, so adding columns cannot break a test."""
    column = [c.value for c in sheet[1]].index(heading) + 1
    return sheet.cell(row=row, column=column)


def test_there_is_a_header_row_and_one_row_per_post(sheet):
    assert sheet.max_row == 3
    assert sheet["A1"].value == "#"


def test_rows_are_numbered_in_rank_order(sheet):
    assert sheet["A2"].value == 1
    assert sheet["A3"].value == 2


def test_the_title_links_to_the_post(sheet):
    target = cell(sheet, 2, "Title").hyperlink.target
    assert target == "https://www.reddit.com/r/ubuntu/comments/1"


def test_metrics_are_written_as_numbers_not_text(sheet):
    points = cell(sheet, 2, "Points").value
    assert points == 900
    assert isinstance(points, int)


def test_missing_metrics_are_left_empty_rather_than_zero(sheet):
    assert cell(sheet, 3, "Points").value is None


def test_the_header_row_is_frozen_so_it_stays_while_scrolling(tmp_path):
    target = tmp_path / "frozen.xlsx"
    write_xlsx(RANKED, target)
    assert openpyxl.load_workbook(target).active.freeze_panes == "A2"


def test_an_empty_run_still_writes_a_usable_sheet(tmp_path):
    target = tmp_path / "empty.xlsx"
    write_xlsx([], target)
    sheet = openpyxl.load_workbook(target).active
    assert sheet.max_row == 1
    assert sheet["A1"].value == "#"


DATED = Post(
    title="Ubuntu 26.04 released",
    subreddit="ubuntu",
    permalink="/r/ubuntu/comments/1",
    author="someone",
    published="2026-09-23T11:23:09+00:00",
)


def test_author_and_date_have_their_own_columns(tmp_path):
    target = tmp_path / "dated.xlsx"
    write_xlsx([DATED], target)
    sheet = openpyxl.load_workbook(target).active
    assert "Author" in [c.value for c in sheet[1]]
    assert "Posted" in [c.value for c in sheet[1]]


def test_the_timestamp_is_shortened_to_something_readable(tmp_path):
    target = tmp_path / "dated.xlsx"
    write_xlsx([DATED], target)
    sheet = openpyxl.load_workbook(target).active
    assert cell(sheet, 2, "Posted").value == "2026-09-23 11:23"


def test_the_post_body_gets_its_own_column(tmp_path):
    target = tmp_path / "body.xlsx"
    written = Post(title="t", subreddit="ubuntu", summary="I have 8GB RAM.")
    write_xlsx([written], target)
    sheet = openpyxl.load_workbook(target).active
    assert cell(sheet, 2, "Body").value == "I have 8GB RAM."


SOURCED = Post(
    title="Hurricane Polo hits Category 5",
    subreddit="news",
    permalink="/r/news/comments/3",
    source_url="https://www.cnn.com/2026/09/22/weather/hurricane-polo",
)


def test_the_source_column_shows_the_domain_not_the_whole_url(tmp_path):
    target = tmp_path / "sourced.xlsx"
    write_xlsx([SOURCED], target)
    sheet = openpyxl.load_workbook(target).active
    assert cell(sheet, 2, "Source").value == "cnn.com"


def test_the_source_cell_links_to_the_article(tmp_path):
    target = tmp_path / "sourced.xlsx"
    write_xlsx([SOURCED], target)
    sheet = openpyxl.load_workbook(target).active
    assert cell(sheet, 2, "Source").hyperlink.target.startswith("https://www.cnn.com/")


def test_a_post_with_no_article_leaves_the_source_cell_empty(tmp_path):
    target = tmp_path / "plain.xlsx"
    write_xlsx([RANKED[0]], target)
    sheet = openpyxl.load_workbook(target).active
    assert cell(sheet, 2, "Source").value is None
