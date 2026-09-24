"""An Excel version of the briefing.

CSV is fine for feeding another script. This is for the person who opens the file,
sorts it, and clicks through to the threads, so the titles are real hyperlinks and
the numbers are numbers rather than text.

openpyxl is an optional dependency: pip install "reddit-trend-radar[xlsx]"
"""

from __future__ import annotations

from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from radar.models import Post

HEADERS = (
    "#", "Subreddit", "Posted", "Author", "Title", "Source", "Body", "Points", "Comments",
)
WIDTHS = (5, 16, 18, 20, 66, 30, 80, 10, 12)


def write_xlsx(ranked: list[Post], path: str | Path) -> None:
    """Write the ranked posts to a spreadsheet with clickable titles."""
    workbook = _new_workbook()
    sheet = workbook.active
    sheet.title = "Trend Radar"
    _write_header(sheet)
    for row, post in enumerate(ranked, start=2):
        _write_post(sheet, row, post)
    _apply_layout(sheet, len(ranked))
    workbook.save(str(path))


def _new_workbook() -> Any:
    try:
        from openpyxl import Workbook
    except ImportError:  # pragma: no cover - exercised only without the extra installed
        raise RuntimeError(
            'Excel output needs openpyxl: pip install "reddit-trend-radar[xlsx]"'
        ) from None
    return Workbook()


def _write_header(sheet: Any) -> None:
    from openpyxl.styles import Font

    for column, heading in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=heading)
        cell.font = Font(bold=True)


def _write_post(sheet: Any, row: int, post: Post) -> None:
    from openpyxl.styles import Alignment, Font

    sheet.cell(row=row, column=1, value=row - 1)
    sheet.cell(row=row, column=2, value=post.subreddit)
    sheet.cell(row=row, column=3, value=post.posted_at or None)
    sheet.cell(row=row, column=4, value=post.author or None)
    title = sheet.cell(row=row, column=5, value=post.title)
    if post.link:
        title.hyperlink = post.link
        title.font = Font(color="0563C1", underline="single")
    source = sheet.cell(row=row, column=6, value=_host(post.source_url) or None)
    if post.source_url:
        source.hyperlink = post.source_url
        source.font = Font(color="0563C1", underline="single")
    body = sheet.cell(row=row, column=7, value=post.summary or None)
    body.alignment = Alignment(wrap_text=True, vertical="top")
    # Left empty rather than zeroed: the Atom feed gives no counts, and a blank
    # cell says that honestly while a 0 would not.
    sheet.cell(row=row, column=8, value=post.score)
    sheet.cell(row=row, column=9, value=post.num_comments)


def _host(url: str) -> str:
    """Just the domain, so the column reads as a source rather than a wall of URL."""
    return urlparse(url).netloc.removeprefix("www.") if url else ""


def _apply_layout(sheet: Any, post_count: int) -> None:
    from openpyxl.utils import get_column_letter

    for column, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    if post_count:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{post_count + 1}"
